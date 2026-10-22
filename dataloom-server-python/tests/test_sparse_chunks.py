"""Run with: python -m unittest discover -s tests -v"""
import json
import tempfile
import unittest
from pathlib import Path

import openpyxl
from sqlalchemy import create_engine
from sqlalchemy.orm import Session

from app.config import CHUNK_SIZE
from app.database import Base
from app.models import ExcelDocument, ExcelSheet
from app.services.parser_service import parse_and_save
from app.services.sheet_service import (
    batch_update_cells, list_chunks_by_sheet_id, replace_workbook,
)


class SparseChunkTests(unittest.TestCase):
    def setUp(self):
        self.engine = create_engine("sqlite://")
        Base.metadata.create_all(self.engine)
        self.db = Session(self.engine)
        self.doc = ExcelDocument(name="sparse.xlsx")
        self.db.add(self.doc)
        self.db.commit()
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.addCleanup(self.engine.dispose)
        self.addCleanup(self.db.close)

    def upload(self, values):
        wb = openpyxl.Workbook()
        for row, value in values.items():
            wb.active.cell(row=row + 1, column=1, value=value)
        path = Path(self.tmp.name) / "sparse.xlsx"
        wb.save(path)
        wb.close()
        sheet = parse_and_save(str(path), self.doc, self.db)[0]
        self.db.commit()
        return sheet

    def snapshot(self, sheet, expected):
        self.db.expire_all()
        chunks = list_chunks_by_sheet_id(sheet.id, self.db)
        indices = [chunk.chunk_index for chunk in chunks]
        self.assertEqual(len(indices), len(set(indices)))
        cells = []
        for chunk in chunks:
            self.assertEqual(chunk.row_start, chunk.chunk_index * CHUNK_SIZE)
            self.assertEqual(chunk.row_end, (chunk.chunk_index + 1) * CHUNK_SIZE - 1)
            for cell in json.loads(chunk.celldata_json):
                self.assertEqual(chunk.chunk_index, cell["r"] // CHUNK_SIZE)
                cells.append(cell)
        coordinates = [(cell["r"], cell["c"]) for cell in cells]
        self.assertEqual(len(coordinates), len(set(coordinates)), "duplicate cells")
        actual = {(cell["r"], cell["c"]): cell["v"] for cell in cells}
        self.assertEqual(actual, expected)
        return cells

    def exercise(self, rows):
        sheet = self.upload(rows)
        chunks = list_chunks_by_sheet_id(sheet.id, self.db)
        expected = {(cell["r"], cell["c"]): cell["v"]
                    for chunk in chunks for cell in json.loads(chunk.celldata_json)}
        self.assertEqual(set(expected), {(r, 0) for r in rows})
        self.snapshot(sheet, expected)
        self.assertEqual(sheet.chunk_count, max(rows) // CHUNK_SIZE + 1)
        self.assertEqual(len(chunks), len({r // CHUNK_SIZE for r in rows}))

        # Edit every uploaded coordinate, create a cell in an omitted block,
        # and send repeated updates to ensure the last value replaces the old one.
        updates = []
        for r in rows:
            expected[(r, 0)] = {"v": f"edited-{r}"}
            updates.extend([{"sheetId": sheet.id, "r": r, "c": 0, "v": {"v": "old"}},
                            {"sheetId": sheet.id, "r": r, "c": 0, "v": expected[(r, 0)]}])
        expected[(0, 1)] = {"f": "=1+1", "v": "", "m": ""}
        updates.append({"sheetId": sheet.id, "r": 0, "c": 1, "v": expected[(0, 1)]})
        batch_update_cells(self.doc.id, updates, self.db)
        cells = self.snapshot(sheet, expected)

        config = {"merge": {}, "columnlen": {"0": 96}, "rowlen": {}}
        mapping = replace_workbook(self.doc.id, [{"name": "Sparse", "row": max(rows) + 1,
            "column": 2, "celldata": cells, "config": config}], self.db)
        sheet = self.db.get(ExcelSheet, mapping[0])
        self.snapshot(sheet, expected)
        self.assertEqual(json.loads(sheet.config_json), config)
        self.assertEqual(sheet.chunk_count, max(rows) // CHUNK_SIZE + 1)
        # Delete across chunk boundaries; empty physical chunks remain compatible.
        batch_update_cells(self.doc.id, [{"sheetId": sheet.id, "r": r, "c": c, "v": None}
                                       for r, c in expected], self.db)
        self.snapshot(sheet, {})

    def test_upload_then_edit_does_not_duplicate_a1001(self):
        sheet = self.upload({CHUNK_SIZE: "original"})
        batch_update_cells(self.doc.id, [{"sheetId": sheet.id, "r": CHUNK_SIZE,
            "c": 0, "v": {"v": "changed"}}], self.db)
        cells = [cell for chunk in list_chunks_by_sheet_id(sheet.id, self.db)
                 for cell in json.loads(chunk.celldata_json)]
        self.assertEqual(cells, [{"r": CHUNK_SIZE, "c": 0, "v": {"v": "changed"}}])

    def test_leading_blank_block(self):
        self.exercise({CHUNK_SIZE: "A1001"})

    def test_multiple_leading_blank_blocks(self):
        self.exercise({2 * CHUNK_SIZE: "late"})

    def test_middle_blank_block(self):
        self.exercise({0: "first", 2 * CHUNK_SIZE: "third"})

    def test_boundaries(self):
        self.exercise({CHUNK_SIZE - 1: 0, CHUNK_SIZE: False,
                       2 * CHUNK_SIZE - 1: "end", 2 * CHUNK_SIZE: "next"})

    def test_uncached_formula_after_blank_block(self):
        sheet = self.upload({CHUNK_SIZE: "=1+1"})
        expected = {(CHUNK_SIZE, 0): {"f": "=1+1", "v": "", "m": "", "ct": {"fa": "General", "t": "s"}}}
        cells = self.snapshot(sheet, expected)
        mapping = replace_workbook(self.doc.id, [{"celldata": cells}], self.db)
        sheet = self.db.get(ExcelSheet, mapping[0])
        self.snapshot(sheet, expected)
        expected[(CHUNK_SIZE, 0)] = {"f": "=2+2", "v": 4}
        batch_update_cells(self.doc.id, [{"sheetId": sheet.id, "r": CHUNK_SIZE, "c": 0,
                                        "v": expected[(CHUNK_SIZE, 0)]}], self.db)
        self.snapshot(sheet, expected)

    def test_new_chunk_extends_index_bound(self):
        sheet = self.upload({0: "first"})
        value = {"v": "later"}
        batch_update_cells(self.doc.id, [{"sheetId": sheet.id, "r": 3 * CHUNK_SIZE,
            "c": 0, "v": value}], self.db)
        self.db.expire_all()
        self.assertEqual(sheet.chunk_count, 4)
        chunks = list_chunks_by_sheet_id(sheet.id, self.db)
        self.assertEqual([chunk.chunk_index for chunk in chunks], [0, 3])
        self.assertEqual(json.loads(chunks[1].celldata_json),
                         [{"r": 3 * CHUNK_SIZE, "c": 0, "v": value}])

    def test_empty_full_save_metadata(self):
        mapping = replace_workbook(self.doc.id, [{"row": 2 * CHUNK_SIZE + 1}], self.db)
        sheet = self.db.get(ExcelSheet, mapping[0])
        self.snapshot(sheet, {})
        self.assertEqual(sheet.chunk_count, 3)

    def test_blank_workbook_and_matrix_full_save(self):
        sheet = self.upload({})
        self.snapshot(sheet, {})
        self.assertEqual(sheet.chunk_count, 1)
        data = [None] * CHUNK_SIZE + [[{"v": "matrix"}]]
        mapping = replace_workbook(self.doc.id, [{"data": data}], self.db)
        sheet = self.db.get(ExcelSheet, mapping[0])
        self.snapshot(sheet, {(CHUNK_SIZE, 0): {"v": "matrix"}})
        self.assertEqual(sheet.chunk_count, 2)


if __name__ == "__main__":
    unittest.main()
